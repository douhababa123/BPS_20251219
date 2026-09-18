from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from annual_baseline import build_baseline_export, build_template, parse_baseline_workbook, preview_payload


MIGRATION = Path(__file__).resolve().parents[1] / "migrations" / "008_capability_governance.sql"
DEPLOY_WORKFLOW = Path(__file__).resolve().parents[2] / ".github" / "workflows" / "jetson-deploy.yml"


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current_Target states"
    sheet.append(["Employee ID", "Skill ID", "Initial Current", "Annual Target"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


EMPLOYEES = {"E001": {"id": "employee-uuid", "name": "Employee One"}}
SKILLS = {7: {"module_id": 2, "module_name": "Module", "skill_name": "Skill"}}


def raw_workbook_bytes(*, sheet_name="Current_Target states", current=1, target=2):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet["A3"] = "BPS elements"
    sheet["C5"] = "BPS System approach"
    sheet["C6"] = "C"
    sheet["D6"] = "T"
    sheet["A7"] = "TEST"
    sheet["B7"] = "Employee One"
    sheet["C7"] = current
    sheet["D7"] = target
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_valid_zero_is_kept_and_missing_cells_are_not_created():
    contents = workbook_bytes([["E001", 7, 0, 3]])
    parsed = parse_baseline_workbook(contents, EMPLOYEES, SKILLS)
    preview = preview_payload(contents, parsed, "baseline.xlsx")

    assert preview["valid"] is True
    assert preview["summary"] == {
        "employeeCount": 1,
        "skillCount": 1,
        "cellCount": 1,
        "initialLevel": 0.0,
        "targetLevel": 3.0,
        "initialGap": 3,
    }
    assert len(parsed["cells"]) == 1


def test_any_invalid_or_duplicate_row_blocks_complete_import():
    contents = workbook_bytes([
        ["E001", 7, 1, 3],
        ["E001", 7, 2, 1],
        ["UNKNOWN", 7, 0, 1],
    ])
    parsed = parse_baseline_workbook(contents, EMPLOYEES, SKILLS)
    preview = preview_payload(contents, parsed, "baseline.xlsx")

    assert preview["valid"] is False
    messages = [error["message"] for error in preview["errors"]]
    assert any("重复" in message for message in messages)
    assert any("不能小于" in message for message in messages)
    assert any("不存在" in message for message in messages)


def test_template_has_blank_scores_instead_of_current_values():
    contents = build_template([("E001", "Employee One", 7, "Module", "Skill")])
    workbook = load_workbook(BytesIO(contents), data_only=True)
    sheet = workbook.active
    assert sheet.title == "Current_Target states"
    assert sheet["A2"].value == "E001"
    assert sheet["F2"].value is None
    assert sheet["G2"].value is None
    workbook.close()


def test_active_baseline_export_contains_values_and_can_be_reimported():
    contents = build_baseline_export([
        ("E001", "Employee One", 7, "Module", "Skill", 1, 3),
    ])
    workbook = load_workbook(BytesIO(contents), data_only=True)
    sheet = workbook["Current_Target states"]
    assert sheet["A2"].value == "E001"
    assert sheet["F2"].value == 1
    assert sheet["G2"].value == 3
    workbook.close()

    parsed = parse_baseline_workbook(contents, EMPLOYEES, SKILLS, year=2026)
    assert parsed["errors"] == []
    assert parsed["cells"][0].initial_current == 1
    assert parsed["cells"][0].annual_target == 3


def test_missing_exact_current_target_sheet_lists_detected_sheets():
    contents = raw_workbook_bytes(sheet_name="Current Target states")

    parsed = parse_baseline_workbook(contents, EMPLOYEES, SKILLS, year=2026)

    assert parsed["cells"] == []
    assert parsed["errors"][0]["field"] == "worksheet"
    assert "Current_Target states" in parsed["errors"][0]["message"]
    assert "Current Target states" in parsed["errors"][0]["message"]


def test_original_wide_sheet_is_converted_with_source_metadata():
    employees = {"E001": {"id": "employee-uuid", "name": "Employee One", "code": "E001"}}
    skills = {
        28: {
            "module_id": 1,
            "module_name": "BPS elements",
            "skill_name": "Leading in a BPS Plant",
        }
    }
    contents = raw_workbook_bytes(current=None, target=2)

    parsed = parse_baseline_workbook(contents, employees, skills, year=2026)
    preview = preview_payload(contents, parsed, "original.xlsx")

    assert preview["valid"] is True
    assert preview["source"] == {
        "sheetName": "Current_Target states",
        "layout": "WIDE_CT",
        "sourceEmployeeCount": 1,
        "omittedCellCount": 0,
        "excludedEmployeeCount": 0,
        "ignoredSkillColumnCount": 0,
        "errorCount": 0,
    }
    assert preview["rows"][0]["sourceCells"] == "C7/D7"
    assert preview["rows"][0]["conversionRule"] == "C 为空，按 0 导入"
    assert preview["rows"][0]["initialCurrent"] == 0
    assert preview["rowsTruncated"] is False


def test_original_material_supply_header_accepts_current_name_with_or_without_closing_parenthesis():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Current_Target states"
    sheet["C6"] = "C"
    sheet["D6"] = "T"
    sheet["AW5"] = "Material supply(POUP, Milkrun AGV)"
    sheet["AW6"] = "C"
    sheet["AX6"] = "T"
    sheet["A7"] = "TEST"
    sheet["B7"] = "Employee One"
    sheet["AW7"] = 1
    sheet["AX7"] = 2
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    employees = {"E001": {"id": "employee-uuid", "name": "Employee One", "code": "E001"}}

    for skill_name in ("Material supply(POUP, Milkrun AGV", "Material supply(POUP, Milkrun AGV)"):
        parsed = parse_baseline_workbook(output.getvalue(), employees, {
            50: {"module_id": 5, "module_name": "Waste-free, stable flow_LBP", "skill_name": skill_name}
        }, year=2026)
        assert parsed["errors"] == []
        assert len(parsed["cells"]) == 1


def test_wide_sheet_blank_pair_is_omitted_not_zero_filled():
    employees = {"E001": {"id": "employee-uuid", "name": "Employee One", "code": "E001"}}
    skills = {
        28: {
            "module_id": 1,
            "module_name": "BPS elements",
            "skill_name": "Leading in a BPS Plant",
        }
    }
    contents = raw_workbook_bytes(current=None, target=None)

    parsed = parse_baseline_workbook(contents, employees, skills, year=2026)
    preview = preview_payload(contents, parsed, "original.xlsx")

    assert preview["valid"] is False
    assert preview["source"]["omittedCellCount"] == 1
    assert any("没有可导入" in error["message"] for error in preview["errors"])


def test_standard_sheet_omits_resource_only_employees():
    contents = workbook_bytes([["TYLER", 7, 2, 3], ["E001", 7, 1, 2]])
    employees = {
        **EMPLOYEES,
        "TYLER": {"id": "tyler-uuid", "name": "Tyler Tan", "code": "TYLER"},
    }

    parsed = parse_baseline_workbook(contents, employees, SKILLS, year=2026)
    preview = preview_payload(contents, parsed, "baseline.xlsx")

    assert preview["valid"] is True
    assert preview["summary"]["employeeCount"] == 1
    assert preview["source"]["excludedEmployeeCount"] == 1


def test_migration_is_additive_and_does_not_promote_july_snapshot():
    sql = MIGRATION.read_text(encoding="utf-8")
    assert "CREATE TABLE dbo.competency_assessment_versions" in sql
    assert "CREATE TABLE dbo.competency_annual_baselines" in sql
    assert "CREATE TABLE dbo.competency_annual_baseline_items" in sql
    assert "MIGRATION_BASELINE" in sql
    assert "EXEC sp_executesql N'\n            ALTER TABLE dbo.competency_assessment_history" in sql
    assert "EXEC sp_executesql N'\n            CREATE INDEX IX_competency_history_version" in sql
    assert "EXEC sp_executesql @drop_history_source_constraint_sql" in sql
    assert "EXEC(N'ALTER TABLE dbo.competency_assessment_history" not in sql
    assert "INSERT INTO dbo.competency_annual_baselines" not in sql
    assert "UPDATE dbo.competency_assessments" not in sql
    assert "DELETE FROM dbo.competency" not in sql


def test_jetson_deploy_applies_capability_governance_schema_before_startup():
    workflow = DEPLOY_WORKFLOW.read_text(encoding="utf-8")

    migration_at = workflow.index("migrations/008_capability_governance.sql")
    startup_at = workflow.index("docker compose up -d", migration_at)

    assert migration_at < startup_at
