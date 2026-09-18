"""Excel parsing and summaries for immutable annual competency baselines."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from typing import Any, Dict, Iterable, List, Mapping, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


MAX_BASELINE_FILE_BYTES = 5 * 1024 * 1024
SOURCE_SHEET_NAME = "Current_Target states"

RAW_MODULE_RANGES = (
    (3, 14, "BPS elements"),
    (15, 24, "Investment efficiency_PGL"),
    (25, 32, "Waste-free, stable flow_IE"),
    (33, 42, "Waste-free, stable flow_TPM"),
    (43, 52, "Waste-free, stable flow_LBP"),
    (53, 58, "Everybody's CIP"),
    (59, 68, "Leadership commitment"),
    (69, 78, "CIP in indirect area_LEAN"),
    (79, 84, "Digital Transformation"),
)

RAW_SKILL_ALIASES = {
    "bps system approach": "leading in a bps plant",
    "business analysis and spec. preparation of dgital product":
        "business analysis and spec. preparation of digital product",
    "material supply(poup, milkrun agv)": "material supply(poup, milkrun agv",
}

RESOURCE_ONLY_EMPLOYEES = {"tyler tan", "tong zhifeng"}

HEADER_ALIASES = {
    "employee_id": {"employee id", "employee_id", "员工id", "员工编号", "员工工号"},
    "skill_id": {"skill id", "skill_id", "技能id", "能力id"},
    "initial_current": {"initial current", "initial_current", "年初现状", "年初level", "年初 level"},
    "annual_target": {"annual target", "annual_target", "年度目标", "目标level", "目标 level"},
}


@dataclass(frozen=True)
class BaselineCell:
    row_number: int
    employee_id: str
    employee_code: str
    employee_name: str
    skill_id: int
    module_id: int
    module_name: str
    skill_name: str
    initial_current: int
    annual_target: int
    source_cells: str = ""
    conversion_rule: str = "直接导入"

    @property
    def gap(self) -> int:
        return self.annual_target - self.initial_current


def file_sha256(contents: bytes) -> str:
    return sha256(contents).hexdigest()


def _normalized_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _normalized_text(value: Any) -> str:
    return " ".join(
        str(value or "")
        .strip()
        .lower()
        .replace("\n", " ")
        .replace("…", "...")
        .replace("’", "'")
        .split()
    )


def _normalized_skill_name(value: Any) -> str:
    normalized = _normalized_text(value)
    return RAW_SKILL_ALIASES.get(normalized, normalized)


def _header_indexes(header_values: Iterable[Any]) -> Tuple[Dict[str, int], List[Dict[str, Any]]]:
    normalized = [_normalized_header(value) for value in header_values]
    indexes: Dict[str, int] = {}
    errors: List[Dict[str, Any]] = []
    for field, aliases in HEADER_ALIASES.items():
        alias_norms = {_normalized_header(alias) for alias in aliases}
        matches = [index for index, value in enumerate(normalized) if value in alias_norms]
        if len(matches) == 1:
            indexes[field] = matches[0]
        elif not matches:
            errors.append({"row": 1, "field": field, "message": "缺少必需列"})
        else:
            errors.append({"row": 1, "field": field, "message": "列名重复"})
    return indexes, errors


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _level(value: Any) -> int | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    if not numeric.is_integer():
        return None
    integer = int(numeric)
    return integer if 0 <= integer <= 4 else None


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _source_metadata(layout: str = "UNKNOWN") -> Dict[str, Any]:
    return {
        "sheetName": SOURCE_SHEET_NAME,
        "layout": layout,
        "sourceEmployeeCount": 0,
        "omittedCellCount": 0,
        "excludedEmployeeCount": 0,
        "ignoredSkillColumnCount": 0,
    }


def _parse_standard_sheet(
    worksheet,
    employees_by_code: Mapping[str, Mapping[str, Any]],
    skills_by_id: Mapping[int, Mapping[str, Any]],
) -> Dict[str, Any]:
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return {"cells": [], "errors": [{"row": 1, "field": "file", "message": "工作表为空"}], "source": _source_metadata("LONG_FORM")}

    indexes, errors = _header_indexes(header)
    if errors:
        return {"cells": [], "errors": errors, "source": _source_metadata("LONG_FORM")}

    cells: List[BaselineCell] = []
    seen: Dict[Tuple[str, int], int] = {}
    source_employees = set()
    excluded_employees = set()
    normalized_employees = {str(key).strip().lower(): value for key, value in employees_by_code.items()}

    for row_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        employee_code = _identifier(values[indexes["employee_id"]] if indexes["employee_id"] < len(values) else None)
        skill_text = _identifier(values[indexes["skill_id"]] if indexes["skill_id"] < len(values) else None)
        current_raw = values[indexes["initial_current"]] if indexes["initial_current"] < len(values) else None
        target_raw = values[indexes["annual_target"]] if indexes["annual_target"] < len(values) else None
        if employee_code:
            source_employees.add(employee_code.lower())

        employee = normalized_employees.get(employee_code.lower())
        if not employee_code:
            errors.append({"row": row_number, "field": "employee_id", "message": "员工 ID 不能为空"})
        elif not employee:
            errors.append({"row": row_number, "field": "employee_id", "message": "员工不存在或未启用"})
        elif _normalized_text(employee.get("name")) in RESOURCE_ONLY_EMPLOYEES:
            excluded_employees.add(str(employee["id"]).lower())
            continue

        try:
            skill_id = int(skill_text)
        except (TypeError, ValueError):
            skill_id = -1
        skill = skills_by_id.get(skill_id)
        if not skill_text:
            errors.append({"row": row_number, "field": "skill_id", "message": "技能 ID 不能为空"})
        elif not skill:
            errors.append({"row": row_number, "field": "skill_id", "message": "技能不存在或未启用"})

        current = _level(current_raw)
        target = _level(target_raw)
        if current is None:
            errors.append({"row": row_number, "field": "initial_current", "message": "年初现状必须是 0–4 的整数"})
        if target is None:
            errors.append({"row": row_number, "field": "annual_target", "message": "年度目标必须是 0–4 的整数"})
        if current is not None and target is not None and target < current:
            errors.append({"row": row_number, "field": "annual_target", "message": "年度目标不能小于年初现状"})

        if employee and skill:
            key = (str(employee["id"]).lower(), skill_id)
            if key in seen:
                errors.append({
                    "row": row_number,
                    "field": "employee_id,skill_id",
                    "message": f"与第 {seen[key]} 行重复",
                })
            else:
                seen[key] = row_number

        if employee and skill and current is not None and target is not None and target >= current and seen.get((str(employee["id"]).lower(), skill_id)) == row_number:
            cells.append(BaselineCell(
                row_number=row_number,
                employee_id=str(employee["id"]),
                employee_code=employee_code,
                employee_name=str(employee.get("name") or ""),
                skill_id=skill_id,
                module_id=int(skill["module_id"]),
                module_name=str(skill.get("module_name") or ""),
                skill_name=str(skill.get("skill_name") or ""),
                initial_current=current,
                annual_target=target,
                source_cells=f"A{row_number}:G{row_number}",
                conversion_rule="标准长表直接导入",
            ))

    source = _source_metadata("LONG_FORM")
    source["sourceEmployeeCount"] = len(source_employees)
    source["excludedEmployeeCount"] = len(excluded_employees)
    return {"cells": cells, "errors": errors, "source": source}


def _raw_module_name(column: int) -> Optional[str]:
    for start, end, name in RAW_MODULE_RANGES:
        if start <= column <= end:
            return name
    return None


def _resolve_raw_skills(worksheet, skills_by_id):
    by_key: Dict[Tuple[str, str], List[Tuple[int, Mapping[str, Any]]]] = {}
    for skill_id, skill in skills_by_id.items():
        key = (_normalized_text(skill.get("module_name")), _normalized_skill_name(skill.get("skill_name")))
        by_key.setdefault(key, []).append((int(skill_id), skill))

    resolved = {}
    ignored = 0
    errors = []
    for column in range(3, 85, 2):
        source_name = _normalized_text(worksheet.cell(5, column).value)
        if not source_name:
            continue
        if source_name in {"...", "bps essential - delete"}:
            ignored += 1
            continue
        module_name = _raw_module_name(column)
        target_name = _normalized_skill_name(source_name)
        matches = by_key.get((_normalized_text(module_name), target_name), [])
        if len(matches) != 1:
            errors.append({
                "row": 5,
                "field": get_column_letter(column),
                "message": f"能力列无法唯一匹配启用技能：{worksheet.cell(5, column).value}",
            })
            continue
        resolved[column] = matches[0]
    return resolved, ignored, errors


def _employee_indexes(employees_by_code):
    by_code = {}
    by_name: Dict[str, List[Mapping[str, Any]]] = {}
    for code, employee in employees_by_code.items():
        item = dict(employee)
        item.setdefault("code", str(code))
        by_code[_normalized_text(code)] = item
        by_code[_normalized_text(item.get("code"))] = item
        by_name.setdefault(_normalized_text(item.get("name")), []).append(item)
    return by_code, by_name


def _parse_raw_sheet(worksheet, employees_by_code, skills_by_id, year: Optional[int]):
    source = _source_metadata("WIDE_CT")
    resolved_skills, ignored, errors = _resolve_raw_skills(worksheet, skills_by_id)
    source["ignoredSkillColumnCount"] = ignored
    by_code, by_name = _employee_indexes(employees_by_code)
    cells: List[BaselineCell] = []
    seen = set()
    employee_rows = []

    for row_number in range(7, worksheet.max_row + 1):
        department = _identifier(worksheet.cell(row_number, 1).value)
        employee_name = _identifier(worksheet.cell(row_number, 2).value)
        if not department and not employee_name:
            if employee_rows:
                break
            continue
        if not employee_name:
            errors.append({"row": row_number, "field": "B", "message": "员工姓名不能为空"})
            continue
        employee_rows.append(row_number)
        normalized_name = _normalized_text(employee_name)
        if normalized_name in RESOURCE_ONLY_EMPLOYEES:
            source["excludedEmployeeCount"] += 1
            continue

        generated_code = f"{department}_{employee_name.replace(' ', '_')}" if department else ""
        employee = by_code.get(_normalized_text(generated_code))
        if not employee:
            name_matches = by_name.get(normalized_name, [])
            employee = name_matches[0] if len(name_matches) == 1 else None
            if len(name_matches) > 1:
                errors.append({"row": row_number, "field": "B", "message": f"员工姓名不唯一：{employee_name}"})
        if not employee:
            if not any(error["row"] == row_number and error["field"] == "B" for error in errors):
                errors.append({"row": row_number, "field": "B", "message": f"员工不存在或未启用：{employee_name}"})
            continue

        for column, (skill_id, skill) in resolved_skills.items():
            current_raw = worksheet.cell(row_number, column).value
            target_raw = worksheet.cell(row_number, column + 1).value
            current_blank = _is_blank(current_raw)
            target_blank = _is_blank(target_raw)
            if current_blank and target_blank:
                source["omittedCellCount"] += 1
                continue

            current = 0 if current_blank else _level(current_raw)
            target = _level(target_raw)
            rule = "C 为空，按 0 导入" if current_blank else "C、T 均有值，直接导入"
            if target_blank:
                is_xu_ship_to_line = (
                    year == 2026
                    and normalized_name == "xu qingyue"
                    and _normalized_text(skill.get("skill_name")) == "ship to line"
                    and current == 1
                )
                if is_xu_ship_to_line:
                    target = 2
                    rule = "Xu Qingyue / Ship to line 临时目标按已确认值 2 导入"
                else:
                    errors.append({
                        "row": row_number,
                        "field": get_column_letter(column + 1),
                        "message": "C 有值但 T 为空，年度目标必须填写",
                    })
                    continue
            if current is None:
                errors.append({"row": row_number, "field": get_column_letter(column), "message": "C 必须是 0–4 的整数"})
                continue
            if target is None:
                errors.append({"row": row_number, "field": get_column_letter(column + 1), "message": "T 必须是 0–4 的整数"})
                continue
            if target < current:
                errors.append({"row": row_number, "field": get_column_letter(column + 1), "message": "T 不能小于 C"})
                continue

            key = (str(employee["id"]).lower(), skill_id)
            if key in seen:
                errors.append({"row": row_number, "field": get_column_letter(column), "message": "员工与技能组合重复"})
                continue
            seen.add(key)
            cells.append(BaselineCell(
                row_number=row_number,
                employee_id=str(employee["id"]),
                employee_code=str(employee.get("code") or generated_code),
                employee_name=str(employee.get("name") or employee_name),
                skill_id=skill_id,
                module_id=int(skill["module_id"]),
                module_name=str(skill.get("module_name") or ""),
                skill_name=str(skill.get("skill_name") or ""),
                initial_current=current,
                annual_target=target,
                source_cells=f"{get_column_letter(column)}{row_number}/{get_column_letter(column + 1)}{row_number}",
                conversion_rule=rule,
            ))

    source["sourceEmployeeCount"] = len(employee_rows)
    return {"cells": cells, "errors": errors, "source": source}


def parse_baseline_workbook(
    contents: bytes,
    employees_by_code: Mapping[str, Mapping[str, Any]],
    skills_by_id: Mapping[int, Mapping[str, Any]],
    year: Optional[int] = None,
) -> Dict[str, Any]:
    """Parse the exact baseline worksheet without writing to a database."""

    if not contents:
        return {"cells": [], "errors": [{"row": 0, "field": "file", "message": "文件为空"}], "source": _source_metadata()}
    if len(contents) > MAX_BASELINE_FILE_BYTES:
        return {"cells": [], "errors": [{"row": 0, "field": "file", "message": "文件超过 5 MB"}], "source": _source_metadata()}

    try:
        workbook = load_workbook(BytesIO(contents), read_only=False, data_only=True)
    except Exception:
        return {"cells": [], "errors": [{"row": 0, "field": "file", "message": "无法解析 .xlsx 文件"}], "source": _source_metadata()}

    try:
        if SOURCE_SHEET_NAME not in workbook.sheetnames:
            detected = "、".join(workbook.sheetnames) or "无"
            return {
                "cells": [],
                "errors": [{
                    "row": 0,
                    "field": "worksheet",
                    "message": f"缺少工作表 {SOURCE_SHEET_NAME}；检测到：{detected}",
                }],
                "source": _source_metadata(),
            }
        worksheet = workbook[SOURCE_SHEET_NAME]
        first_row = [cell.value for cell in worksheet[1]]
        standard_indexes, standard_errors = _header_indexes(first_row)
        if not standard_errors and len(standard_indexes) == len(HEADER_ALIASES):
            parsed = _parse_standard_sheet(worksheet, employees_by_code, skills_by_id)
        elif _normalized_text(worksheet.cell(6, 3).value) == "c" and _normalized_text(worksheet.cell(6, 4).value) == "t":
            parsed = _parse_raw_sheet(worksheet, employees_by_code, skills_by_id, year)
        else:
            parsed = {
                "cells": [],
                "errors": [{"row": 1, "field": "worksheet", "message": "无法识别工作表格式：既不是标准长表，也不是原始 C/T 宽表"}],
                "source": _source_metadata(),
            }
        if not parsed["cells"] and not parsed["errors"]:
            parsed["errors"].append({"row": 0, "field": "file", "message": "没有可导入的基线明细"})
        return parsed
    finally:
        workbook.close()


def preview_payload(contents: bytes, parsed: Dict[str, Any], filename: str) -> Dict[str, Any]:
    cells: List[BaselineCell] = parsed["cells"]
    count = len(cells)
    total_current = sum(cell.initial_current for cell in cells)
    total_target = sum(cell.annual_target for cell in cells)
    return {
        "filename": filename,
        "sha256": file_sha256(contents),
        "valid": not parsed["errors"] and count > 0,
        "errors": parsed["errors"],
        "source": {
            **parsed.get("source", _source_metadata()),
            "errorCount": len(parsed["errors"]),
        },
        "summary": {
            "employeeCount": len({cell.employee_id for cell in cells}),
            "skillCount": len({cell.skill_id for cell in cells}),
            "cellCount": count,
            "initialLevel": round(total_current / count, 1) if count else None,
            "targetLevel": round(total_target / count, 1) if count else None,
            "initialGap": total_target - total_current,
        },
        "rows": [
            {
                "row": cell.row_number,
                "employeeId": cell.employee_code,
                "employeeName": cell.employee_name,
                "skillId": cell.skill_id,
                "moduleId": cell.module_id,
                "moduleName": cell.module_name,
                "skillName": cell.skill_name,
                "initialCurrent": cell.initial_current,
                "annualTarget": cell.annual_target,
                "gap": cell.gap,
                "sourceCells": cell.source_cells,
                "conversionRule": cell.conversion_rule,
            }
            for cell in cells
        ],
        "rowsTruncated": False,
    }


def build_template(rows: Iterable[Tuple[str, str, int, str, str]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SOURCE_SHEET_NAME
    worksheet.append([
        "Employee ID",
        "Employee Name",
        "Skill ID",
        "Module",
        "Skill",
        "Initial Current",
        "Annual Target",
    ])
    for employee_code, employee_name, skill_id, module_name, skill_name in rows:
        worksheet.append([employee_code, employee_name, skill_id, module_name, skill_name, None, None])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column, width in {"A": 18, "B": 24, "C": 12, "D": 34, "E": 50, "F": 18, "G": 18}.items():
        worksheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
